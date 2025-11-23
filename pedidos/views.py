from django.shortcuts import render, redirect, get_object_or_404
from django.contrib import messages
from django.http import HttpResponse
from productos.models import Producto, Insumo
from .models import Pedido, DetallePedido
import json
from datetime import datetime, timedelta

# Librerías para exportación
from reportlab.lib.pagesizes import letter, A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from io import BytesIO


# HU-01: Registrar nuevo pedido
def registrar_pedido_vista(request):
    if request.method == 'POST':
        pedido_data = request.POST.get('pedido_data')

        if pedido_data:
            try:
                pedido_json = json.loads(pedido_data)
                cliente = request.POST.get('cliente', 'Cliente General')

                # 1. Calcular total
                total = 0
                for item_id, item in pedido_json.items():
                    precio_base = float(item['precioBase'])
                    cantidad = int(item['cantidad'])
                    precio_extras = sum(float(extra['precio']) for extra in item.get('extras', []))
                    subtotal = (precio_base + precio_extras) * cantidad
                    total += subtotal

                # 2. Crear Pedido
                pedido = Pedido.objects.create(
                    cliente=cliente,
                    total=total,
                    estado='En preparación'
                )

                # 3. Crear Detalles
                for item_id, item in pedido_json.items():
                    producto_id = item['pId']
                    cantidad = item['cantidad']

                    lista_extras = [f"{extra['nombre']} (S/{extra['precio']})" for extra in item.get('extras', [])]
                    texto_personalizacion = ", ".join(lista_extras)

                    try:
                        producto = Producto.objects.get(id=producto_id)
                        DetallePedido.objects.create(
                            pedido=pedido,
                            producto=producto,
                            cantidad=cantidad,
                            personalizacion=texto_personalizacion
                        )
                    except Producto.DoesNotExist:
                        pass

                messages.success(request, f'Pedido #{pedido.id} registrado exitosamente!')
                return redirect('panel-pedidos')

            except Exception as e:
                messages.error(request, f'Error al procesar: {str(e)}')
        else:
            messages.warning(request, 'El carrito está vacío')

    productos_bd = Producto.objects.prefetch_related('insumos').all()
    insumos_bd = Insumo.objects.all()

    extras_list = []
    for insumo in insumos_bd:
        extras_list.append({
            'id': insumo.id,
            'nombre': insumo.nombre,
            'precio': float(insumo.precio)
        })

    extras_json = json.dumps(extras_list)

    contexto = {
        'productos': productos_bd,
        'extras_json': extras_json
    }
    return render(request, 'pedidos/registrar_pedido.html', contexto)


# HU-02: Personalizar producto
def personalizar_producto_vista(request, producto_id):
    producto = Producto.objects.get(id=producto_id)
    contexto = {
        'producto': producto
    }
    return render(request, 'pedidos/personalizar_producto.html', contexto)


# HU-03: Panel de pedidos del cajero
def panel_pedidos_vista(request):
    pedidos_activos = Pedido.objects.exclude(estado__in=['Entregado', 'Cancelado']).order_by('-fecha')
    contexto = {
        'pedidos': pedidos_activos
    }
    return render(request, 'pedidos/panel_pedidos.html', contexto)


# Actualizar estado de un pedido
def actualizar_estado_vista(request, pedido_id):
    pedido = get_object_or_404(Pedido, id=pedido_id)

    if request.method == 'POST':
        nuevo_estado = request.POST.get('estado')
        estados_validos = ['En preparación', 'Listo para entregar', 'Entregado', 'Cancelado']
        if nuevo_estado in estados_validos:
            pedido.estado = nuevo_estado
            pedido.save()
            messages.success(request, f'✅ Estado del pedido #{pedido.id} actualizado a: {nuevo_estado}')

    # Verificar de dónde viene la petición
    referer = request.META.get('HTTP_REFERER', '')
    if 'panel-admin' in referer:
        return redirect('panel-pedidos-admin')
    else:
        return redirect('panel-pedidos')


# HU-04: Consultar historial de ventas
def historial_ventas_vista(request):
    fecha_inicio = request.GET.get('fecha_inicio')
    fecha_fin = request.GET.get('fecha_fin')

    ventas = Pedido.objects.filter(estado='Entregado').order_by('-fecha')

    # Validar que se haya seleccionado al menos una fecha al intentar filtrar
    if request.GET and not fecha_inicio and not fecha_fin:
        messages.warning(request, 'Debe seleccionar un periodo para continuar.')

    # Solo filtrar si hay fechas proporcionadas
    if fecha_inicio:
        try:
            fecha_inicio_dt = datetime.strptime(fecha_inicio, '%Y-%m-%d')
            ventas = ventas.filter(fecha__gte=fecha_inicio_dt)
        except ValueError:
            messages.error(request, 'Formato de fecha inicio inválido.')

    if fecha_fin:
        try:
            fecha_fin_dt = datetime.strptime(fecha_fin, '%Y-%m-%d')
            fecha_fin_dt = fecha_fin_dt + timedelta(days=1)
            ventas = ventas.filter(fecha__lt=fecha_fin_dt)
        except ValueError:
            messages.error(request, 'Formato de fecha fin inválido.')

    total_ventas = sum(venta.total for venta in ventas)
    cantidad_ventas = ventas.count()

    contexto = {
        'ventas': ventas,
        'total_ventas': total_ventas,
        'cantidad_ventas': cantidad_ventas,
        'fecha_inicio': fecha_inicio or '',
        'fecha_fin': fecha_fin or '',
    }

    return render(request, 'pedidos/historial_ventas.html', contexto)


def detalle_venta_vista(request, pedido_id):
    venta = get_object_or_404(Pedido, id=pedido_id)
    detalles = venta.detalles.all()

    contexto = {
        'venta': venta,
        'detalles': detalles,
    }

    return render(request, 'pedidos/detalle_venta.html', contexto)


# EXPORTAR A PDF
def exportar_ventas_pdf(request):
    fecha_inicio = request.GET.get('fecha_inicio', '')
    fecha_fin = request.GET.get('fecha_fin', '')

    # Validación: verificar que se hayan seleccionado fechas
    if not fecha_inicio and not fecha_fin:
        messages.warning(request, 'Debe seleccionar un periodo (fecha inicio y/o fecha fin) para generar el reporte.')
        return redirect('historial-ventas')

    ventas = Pedido.objects.filter(estado='Entregado').order_by('-fecha')

    if fecha_inicio:
        try:
            fecha_inicio_dt = datetime.strptime(fecha_inicio, '%Y-%m-%d')
            ventas = ventas.filter(fecha__gte=fecha_inicio_dt)
        except ValueError:
            messages.error(request, 'Formato de fecha inicio inválido.')
            return redirect('historial-ventas')

    if fecha_fin:
        try:
            fecha_fin_dt = datetime.strptime(fecha_fin, '%Y-%m-%d')
            fecha_fin_dt = fecha_fin_dt + timedelta(days=1)
            ventas = ventas.filter(fecha__lt=fecha_fin_dt)
        except ValueError:
            messages.error(request, 'Formato de fecha fin inválido.')
            return redirect('historial-ventas')

    # Validación: verificar que existan ventas en el periodo
    if not ventas.exists():
        messages.warning(request, 'No hay datos para generar el reporte en el periodo seleccionado.')
        url_redirect = f'/historial-ventas/?fecha_inicio={fecha_inicio}&fecha_fin={fecha_fin}'
        return redirect(url_redirect)

    # Crear el PDF
    buffer = BytesIO()
    pdf = SimpleDocTemplate(buffer, pagesize=A4)
    elementos = []

    # Estilos
    estilos = getSampleStyleSheet()
    estilo_titulo = ParagraphStyle(
        'CustomTitle',
        parent=estilos['Heading1'],
        fontSize=18,
        textColor=colors.HexColor('#722f37'),
        spaceAfter=30,
        alignment=TA_CENTER
    )

    # Título
    titulo = Paragraph("Reporte de Ventas - Coco's Burger", estilo_titulo)
    elementos.append(titulo)

    # Periodo
    periodo_texto = f"Periodo: "
    if fecha_inicio:
        periodo_texto += f"{fecha_inicio} "
    if fecha_fin:
        periodo_texto += f"al {fecha_fin}"
    if not fecha_inicio and not fecha_fin:
        periodo_texto += "Todas las ventas"

    periodo = Paragraph(periodo_texto, estilos['Normal'])
    elementos.append(periodo)
    elementos.append(Spacer(1, 0.3 * inch))

    # Resumen
    total_ventas = sum(venta.total for venta in ventas)
    cantidad_ventas = ventas.count()

    resumen_data = [
        ['Total de Ventas:', str(cantidad_ventas)],
        ['Ingresos Totales:', f'S/ {total_ventas:.2f}']
    ]

    resumen_tabla = Table(resumen_data, colWidths=[3 * inch, 2 * inch])
    resumen_tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8f9fa')),
        ('TEXTCOLOR', (0, 0), (-1, -1), colors.black),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 12),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
        ('GRID', (0, 0), (-1, -1), 1, colors.grey)
    ]))
    elementos.append(resumen_tabla)
    elementos.append(Spacer(1, 0.3 * inch))

    # Tabla de ventas
    datos = [['N° Venta', 'Fecha', 'Cliente', 'Monto Total']]

    for venta in ventas:
        datos.append([
            f'#{venta.id}',
            venta.fecha.strftime('%d/%m/%Y %H:%M'),
            venta.cliente,
            f'S/ {venta.total:.2f}'
        ])

    tabla = Table(datos, colWidths=[1 * inch, 1.8 * inch, 2 * inch, 1.5 * inch])
    tabla.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#722f37')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 11),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
        ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
        ('GRID', (0, 0), (-1, -1), 1, colors.black),
        ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
    ]))
    elementos.append(tabla)

    # Generar PDF
    pdf.build(elementos)
    buffer.seek(0)

    # Respuesta HTTP
    response = HttpResponse(buffer, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="reporte_ventas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.pdf"'

    return response


# EXPORTAR A EXCEL
def exportar_ventas_excel(request):
    fecha_inicio = request.GET.get('fecha_inicio', '')
    fecha_fin = request.GET.get('fecha_fin', '')

    # Validación: verificar que se hayan seleccionado fechas
    if not fecha_inicio and not fecha_fin:
        messages.warning(request, 'Debe seleccionar un periodo (fecha inicio y/o fecha fin) para generar el reporte.')
        return redirect('historial-ventas')

    ventas = Pedido.objects.filter(estado='Entregado').order_by('-fecha')

    if fecha_inicio:
        try:
            fecha_inicio_dt = datetime.strptime(fecha_inicio, '%Y-%m-%d')
            ventas = ventas.filter(fecha__gte=fecha_inicio_dt)
        except ValueError:
            messages.error(request, 'Formato de fecha inicio inválido.')
            return redirect('historial-ventas')

    if fecha_fin:
        try:
            fecha_fin_dt = datetime.strptime(fecha_fin, '%Y-%m-%d')
            fecha_fin_dt = fecha_fin_dt + timedelta(days=1)
            ventas = ventas.filter(fecha__lt=fecha_fin_dt)
        except ValueError:
            messages.error(request, 'Formato de fecha fin inválido.')
            return redirect('historial-ventas')

    # Validación: verificar que existan ventas en el periodo
    if not ventas.exists():
        messages.warning(request, 'No hay datos para generar el reporte en el periodo seleccionado.')
        url_redirect = f'/historial-ventas/?fecha_inicio={fecha_inicio}&fecha_fin={fecha_fin}'
        return redirect(url_redirect)

    # Crear el libro de Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte de Ventas"

    # Estilos
    header_fill = PatternFill(start_color="722f37", end_color="722f37", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=12)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # Título
    ws['A1'] = "Reporte de Ventas - Coco's Burger"
    ws['A1'].font = Font(bold=True, size=16, color="722f37")
    ws['A1'].alignment = Alignment(horizontal='center')
    ws.merge_cells('A1:D1')

    # Periodo
    periodo_texto = "Periodo: "
    if fecha_inicio:
        periodo_texto += f"{fecha_inicio} "
    if fecha_fin:
        periodo_texto += f"al {fecha_fin}"
    if not fecha_inicio and not fecha_fin:
        periodo_texto += "Todas las ventas"

    ws['A2'] = periodo_texto
    ws['A2'].font = Font(size=11)
    ws.merge_cells('A2:D2')

    # Resumen
    total_ventas = sum(venta.total for venta in ventas)
    cantidad_ventas = ventas.count()

    ws['A4'] = "Total de Ventas:"
    ws['B4'] = cantidad_ventas
    ws['A5'] = "Ingresos Totales:"
    ws['B5'] = f"S/ {total_ventas:.2f}"

    for row in [4, 5]:
        ws[f'A{row}'].font = Font(bold=True)
        ws[f'A{row}'].fill = PatternFill(start_color="f8f9fa", end_color="f8f9fa", fill_type="solid")
        ws[f'B{row}'].fill = PatternFill(start_color="f8f9fa", end_color="f8f9fa", fill_type="solid")

    # Encabezados de la tabla
    encabezados = ['N° Venta', 'Fecha', 'Cliente', 'Monto Total']
    ws.append([])  # Fila vacía
    ws.append(encabezados)

    header_row = ws[7]
    for cell in header_row:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = border

    # Datos
    for venta in ventas:
        ws.append([
            f'#{venta.id}',
            venta.fecha.strftime('%d/%m/%Y %H:%M'),
            venta.cliente,
            f'S/ {venta.total:.2f}'
        ])

    # Aplicar bordes y alineación a los datos
    for row in ws.iter_rows(min_row=8, max_row=ws.max_row, min_col=1, max_col=4):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # Ajustar ancho de columnas
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 25
    ws.column_dimensions['D'].width = 15

    # Guardar en buffer
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    # Respuesta HTTP
    response = HttpResponse(
        buffer,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="reporte_ventas_{datetime.now().strftime("%Y%m%d_%H%M%S")}.xlsx"'

    return response


# ============================================
# PANEL DE ADMINISTRADOR
# ============================================

# Panel de pedidos para administrador
def panel_pedidos_admin_vista(request):
    from django.utils import timezone

    # Obtener pedidos SOLO del día actual (incluyendo entregados)
    hoy_inicio = timezone.now().replace(hour=0, minute=0, second=0, microsecond=0)
    hoy_fin = hoy_inicio + timedelta(days=1)

    pedidos_hoy = Pedido.objects.filter(
        fecha__gte=hoy_inicio,
        fecha__lt=hoy_fin
    ).order_by('-fecha')

    # Filtro por estado
    estado_filtro = request.GET.get('estado', '')
    if estado_filtro:
        pedidos_hoy = pedidos_hoy.filter(estado=estado_filtro)

    # Estadísticas DIARIAS
    total_pedidos_hoy = pedidos_hoy.count()

    # Ingresos SOLO de pedidos ENTREGADOS
    pedidos_entregados_hoy = pedidos_hoy.filter(estado='Entregado')
    total_ingresos_hoy = sum(pedido.total for pedido in pedidos_entregados_hoy)

    pedidos_preparacion = pedidos_hoy.filter(estado='En preparación').count()
    pedidos_listos = pedidos_hoy.filter(estado='Listo para entregar').count()
    pedidos_entregados = pedidos_entregados_hoy.count()
    pedidos_cancelados = pedidos_hoy.filter(estado='Cancelado').count()

    contexto = {
        'pedidos': pedidos_hoy,
        'total_pedidos_hoy': total_pedidos_hoy,
        'total_ingresos_hoy': total_ingresos_hoy,
        'pedidos_preparacion': pedidos_preparacion,
        'pedidos_listos': pedidos_listos,
        'pedidos_entregados': pedidos_entregados,
        'pedidos_cancelados': pedidos_cancelados,
        'estado_filtro': estado_filtro,
        'fecha_hoy': timezone.now().strftime('%d/%m/%Y'),
    }

    return render(request, 'pedidos/panel_pedidos_admin.html', contexto)


# Finalizar día - ACTUALIZADO
def finalizar_dia_vista(request):
    from django.utils import timezone

    if request.method == 'POST':
        # Obtener pedidos del día actual
        hoy_inicio = timezone.now().replace(hour=0, minute=0, second=0, microsecond=0)
        hoy_fin = hoy_inicio + timedelta(days=1)

        # Pedidos que ya están entregados (se quedan en historial)
        pedidos_entregados = Pedido.objects.filter(
            fecha__gte=hoy_inicio,
            fecha__lt=hoy_fin,
            estado='Entregado'
        )

        # Pedidos activos que NO fueron entregados (se cancelarán)
        pedidos_activos = Pedido.objects.filter(
            fecha__gte=hoy_inicio,
            fecha__lt=hoy_fin,
            estado__in=['En preparación', 'Listo para entregar']
        )

        cantidad_entregados = pedidos_entregados.count()
        cantidad_cancelados = pedidos_activos.count()

        # Cancelar todos los pedidos activos que no se entregaron
        pedidos_activos.update(estado='Cancelado')

        if cantidad_cancelados > 0:
            messages.success(
                request,
                f'✅ Día finalizado. {cantidad_entregados} pedidos entregados en el historial. {cantidad_cancelados} pedidos pendientes fueron cancelados.'
            )
        else:
            messages.success(
                request,
                f'✅ Día finalizado correctamente. {cantidad_entregados} pedidos entregados disponibles en el historial de ventas.'
            )

        return redirect('panel-pedidos-admin')

    return redirect('panel-pedidos-admin')